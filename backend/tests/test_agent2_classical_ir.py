from __future__ import annotations

import asyncio
from pathlib import Path

import pytest
from docx import Document

from backend.app.orchestrator.adapters import build_retrieval_request
from backend.app.orchestrator.agent_clients import (
    LocalRetrievalClient,
    RetrievalClient,
)
from backend.app.retrieval.document_ingestion import (
    DocumentIngestionError,
    DocumentIngestor,
    UnsupportedDocumentError,
    extract_document,
)
from backend.app.retrieval.knowledge_retriever import KnowledgeRetriever
from backend.app.retrieval.preprocessing import (
    preprocess_for_retrieval,
    tokenize_for_retrieval,
)
from backend.app.retrieval.schemas import (
    IntentContext,
    KnowledgeChunk,
    RetrievalRequest,
    RetrievalResponse,
    RetrievalResult,
    UserContext,
)
from backend.app.retrieval.service import RetrievalService
from backend.app.schemas.intake import (
    IntakeData,
    IntakeResponse,
    IntentResult,
)


class MemoryKnowledgeRepository:
    def __init__(self, chunks=None):
        self.chunks = list(chunks or [])

    def list_knowledge_chunks(self, *, insurance_type="motor", document_type=None):
        return [
            item for item in self.chunks
            if item.insurance_type == insurance_type
            and (document_type is None or item.document_type == document_type)
        ]

    def get_knowledge_chunks_by_source(self, source_document_id):
        return [item for item in self.chunks if item.source_document_id == source_document_id]

    def replace_knowledge_chunks(self, source_document_id, chunks):
        self.chunks = [
            item for item in self.chunks if item.source_document_id != source_document_id
        ] + list(chunks)
        return list(chunks)


def _chunk(chunk_id, title, section, content, document_type="policy_manual"):
    return KnowledgeChunk(
        chunk_id=chunk_id,
        source_document_id=f"source-{chunk_id}",
        source_title=title,
        document_type=document_type,
        section=section,
        content=content,
        normalized_content=preprocess_for_retrieval(content),
        metadata={"synthetic": True},
    )


def _corpus():
    return [
        _chunk("flood", "Motor Policy", "Flood", "Flood water damage coverage and policy exclusions for a motor vehicle."),
        _chunk("glass", "Windscreen Guide", "Glass", "Windscreen crack claims require a clear damage photograph."),
        _chunk("theft", "Theft Procedure", "Documents", "A theft claim requires a police report, vehicle registration, and claim form.", "procedure_guide"),
        _chunk("collision", "Collision Guide", "Reporting", "Report a vehicle collision promptly and submit repair estimates.", "guideline"),
    ]


def _minimal_pdf(text: str) -> bytes:
    safe = text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(safe) + 35} >>\nstream\nBT /F1 12 Tf 72 720 Td ({safe}) Tj ET\nendstream",
    ]
    data = b"%PDF-1.4\n"
    offsets = [0]
    for number, body in enumerate(objects, 1):
        offsets.append(len(data))
        data += f"{number} 0 obj\n{body}\nendobj\n".encode()
    xref = len(data)
    data += f"xref\n0 {len(objects)+1}\n0000000000 65535 f \n".encode()
    data += b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:])
    data += f"trailer << /Size {len(objects)+1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode()
    return data


def test_retrieval_preprocessing_preserves_insurance_terms_and_identifiers():
    value = preprocess_for_retrieval("The POLICY MTR-100/26 covers Section 4.2 on flood damage.")
    assert "policy" in value and "mtr-100/26" in value and "4.2" in value
    assert "flood" in value and "damage" in value
    assert "the" not in tokenize_for_retrieval("the vehicle")
    assert preprocess_for_retrieval("   ") == ""
    with pytest.raises(TypeError):
        preprocess_for_retrieval(None)


def test_txt_ingestion_sections_chunks_and_duplicate_protection(tmp_path):
    path = tmp_path / "motor-guide.txt"
    path.write_text("FLOOD COVERAGE:\nFlood water damage guidance.\n\nEXCLUSIONS:\nWear and tear is excluded.", encoding="utf-8")
    repository = MemoryKnowledgeRepository()
    ingestor = DocumentIngestor(repository, max_words=6, overlap_words=1)
    first = ingestor.ingest_file(path, document_type="policy_manual")
    second = ingestor.ingest_file(path, document_type="policy_manual")
    assert first.chunks_created_or_updated >= 2
    assert second.duplicate_unchanged is True
    assert len(repository.chunks) == first.chunks_created_or_updated
    assert {item.section for item in repository.chunks} >= {"FLOOD COVERAGE", "EXCLUSIONS"}
    assert all(item.metadata["content_hash"] for item in repository.chunks)


def test_pdf_and_docx_text_extraction(tmp_path):
    pdf = tmp_path / "flood.pdf"
    pdf.write_bytes(_minimal_pdf("Flood coverage applies according to policy wording."))
    assert "Flood coverage" in extract_document(pdf)[0].content

    docx = tmp_path / "theft.docx"
    document = Document()
    document.add_heading("Theft Documents", level=1)
    document.add_paragraph("Provide a police report and vehicle registration.")
    document.save(docx)
    sections = extract_document(docx)
    assert sections[0].title == "Theft Documents"
    assert "police report" in sections[0].content


def test_unsupported_and_malformed_documents_are_controlled(tmp_path):
    unsupported = tmp_path / "notes.csv"
    unsupported.write_text("not supported")
    with pytest.raises(UnsupportedDocumentError):
        extract_document(unsupported)
    malformed = tmp_path / "bad.pdf"
    malformed.write_bytes(b"not a pdf")
    with pytest.raises(DocumentIngestionError, match="Could not extract"):
        extract_document(malformed)


def test_production_tfidf_ranking_threshold_top_k_and_filter():
    retriever = KnowledgeRetriever(repository=MemoryKnowledgeRepository(_corpus()))
    results = retriever.retrieve("documents required after my car was stolen", top_k=2)
    assert results[0].evidence_id == "theft"
    assert len(results) <= 2
    assert results == sorted(results, key=lambda item: item.score, reverse=True)
    assert results[0].metadata["source_document_id"] == "source-theft"
    filtered = retriever.retrieve(
        "theft claim police report", document_type="procedure_guide"
    )
    assert filtered and all(item.metadata["document_type"] == "procedure_guide" for item in filtered)
    assert retriever.retrieve("quantum spacecraft biology", min_relevance_score=0.08) == []
    assert retriever.retrieve("flood", min_relevance_score=1.0) == []


class EmptyStructuredRepository:
    def get_policy_by_id(self, *args, **kwargs): return None
    def get_policy_by_number(self, *args, **kwargs): return None
    def get_claim_by_id(self, *args, **kwargs): return None
    def get_claim_by_reference(self, *args, **kwargs): return None
    def get_policy_claim_history(self, *args, **kwargs): return []
    def get_claim_documents(self, *args, **kwargs): return []


def test_service_uses_natural_query_and_empty_knowledge_is_no_results():
    retriever = KnowledgeRetriever(repository=MemoryKnowledgeRepository(_corpus()))
    service = RetrievalService(EmptyStructuredRepository(), retriever)
    request = RetrievalRequest(
        request_id="REQ-IR", query="What documents are required for theft?",
        user_context=UserContext(user_id="USR-1"),
        intent_context=IntentContext(intent="required_documents_question"),
    )
    response = service.retrieve(request)
    assert response.status == "success"
    assert response.result.knowledge_evidence[0].evidence_id == "theft"

    empty = RetrievalService(
        EmptyStructuredRepository(),
        KnowledgeRetriever(repository=MemoryKnowledgeRepository()),
    ).retrieve(request)
    assert empty.status == "no_results"
    assert "knowledge_evidence_not_found" in empty.result.missing_evidence


@pytest.mark.parametrize(
    ("intent", "query", "expected_id"),
    [
        ("policy_question", "How do windscreen claims work?", "glass"),
        ("general_information", "How should I report a collision?", "collision"),
    ],
)
def test_service_policy_and_general_information_search_real_corpus(
    intent, query, expected_id
):
    service = RetrievalService(
        EmptyStructuredRepository(),
        KnowledgeRetriever(repository=MemoryKnowledgeRepository(_corpus())),
    )
    response = service.retrieve(
        RetrievalRequest(
            request_id="REQ-INFO", query=query,
            user_context=UserContext(user_id="USR-1"),
            intent_context=IntentContext(intent=intent),
        )
    )
    assert response.status == "success"
    assert response.result.knowledge_evidence[0].evidence_id == expected_id


def test_service_returns_partial_for_generic_evidence_without_requested_policy():
    service = RetrievalService(
        EmptyStructuredRepository(),
        KnowledgeRetriever(repository=MemoryKnowledgeRepository(_corpus())),
    )
    request = RetrievalRequest(
        request_id="REQ-COVER", query="Does my policy cover flood damage?",
        user_context=UserContext(user_id="USR-1"),
        intent_context=IntentContext(intent="coverage_question"),
    )
    response = service.retrieve(request)
    assert response.status == "partial_success"
    assert response.result.knowledge_evidence
    assert "specific_policy_not_available" in response.result.missing_evidence


def test_repository_failure_is_controlled_and_client_is_async_compatible():
    class FailingRetriever:
        def retrieve(self, *args, **kwargs):
            raise RuntimeError("database details that must not leak")

    request = RetrievalRequest(
        request_id="REQ-FAIL", query="motor information",
        user_context=UserContext(user_id="USR-1"),
        intent_context=IntentContext(intent="general_information"),
    )
    service = RetrievalService(EmptyStructuredRepository(), FailingRetriever())
    response = service.retrieve(request)
    assert response.status == "failed"
    assert "database details" not in response.errors[0].message

    client = LocalRetrievalClient(service=service)
    assert isinstance(client, RetrievalClient)
    async_response = asyncio.run(client.retrieve(request))
    assert async_response.status == "failed"


def test_knowledge_repository_exception_is_failed_not_no_results():
    class FailingKnowledgeRepository:
        def list_knowledge_chunks(self, **_kwargs):
            raise RuntimeError("private repository connection detail")

    request = RetrievalRequest(
        request_id="REQ-REPOSITORY-FAIL",
        query="Does my policy cover flood damage?",
        user_context=UserContext(user_id="USR-1"),
        intent_context=IntentContext(intent="coverage_question"),
    )
    response = RetrievalService(
        EmptyStructuredRepository(),
        KnowledgeRetriever(repository=FailingKnowledgeRepository()),
    ).retrieve(request)

    assert response.status == "failed"
    assert response.errors[0].error_code == "knowledge_retrieval_failed"
    assert "private repository" not in response.model_dump_json()


def test_knowledge_failure_with_structured_policy_is_partial_success():
    class PolicyRepository(EmptyStructuredRepository):
        def get_policy_by_number(self, policy_number, user_id):
            return {
                "policy_id": "POL-1", "policy_number": policy_number,
                "customer_id": user_id, "status": "active",
                "start_date": "2026-01-01", "end_date": "2026-12-31",
                "coverage_details": {},
            }

    class FailingRetriever:
        def retrieve(self, *args, **kwargs):
            raise RuntimeError("private database failure")

    from backend.app.retrieval.schemas import PolicyLookupContext

    response = RetrievalService(PolicyRepository(), FailingRetriever()).retrieve(
        RetrievalRequest(
            request_id="REQ-PARTIAL", query="What does my policy cover?",
            user_context=UserContext(user_id="USR-1"),
            intent_context=IntentContext(intent="policy_question"),
            policy_context=PolicyLookupContext(policy_number="MTR-1"),
        )
    )
    assert response.status == "partial_success"
    assert response.result.policy_data is not None
    assert response.errors[0].error_code == "knowledge_retrieval_failed"
    assert "private database" not in response.errors[0].message


def test_orchestrator_adapter_preserves_query_and_absent_identifiers():
    intake = IntakeResponse(
        request_id="REQ-A", status="success",
        data=IntakeData(
            intent=IntentResult(label="coverage_question", confidence=0.9)
        ),
    )
    request = build_retrieval_request(
        request_id="REQ-A", authenticated_user_id="USR-1",
        original_query="Does my policy cover flood damage?", intake=intake,
    )
    assert request.query == "Does my policy cover flood damage?"
    assert request.policy_context is None
    assert request.claim_lookup is None
